"""Rebuild the freedom-spectrum figure (assets/img/posts/constrained-creativity/freedom-spectrum.png) as an
editable PowerPoint slide. Coordinates are measured in pixels on the 688x326 source and mapped 1 px = 0.75 pt."""
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

K = 0.75
FONT, GILL = "Gill Sans", "/System/Library/Fonts/Supplemental/GillSans.ttc"
EDGE = RGBColor(0x20, 0x35, 0x42)
LABEL_PX, AXIS_PX = 19.5, 22

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(688 * K), Pt(326 * K)
s = prs.slides.add_slide(prs.slide_layouts[6])
P = lambda v: Pt(v * K)

# nested regions, drawn back to front: (label, right edge, top edge, fill); all share left = 56, bottom = 255
for name, right, top, fill in [("(4) arbitrary analysis", 594, 28, "9C9C9C"),
                               ("(3) flexible playbooks + operators, fixed DSL", 517, 85, "C1C1C1"),
                               ("(2) flexible playbooks, fixed operators + DSL", 440, 141, "E2E2E2"),
                               ("(1) fixed playbooks", 363, 198, "FFFFFF")]:
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, P(56), P(top), P(right - 56), P(255 - top))
    r.name = name
    r.fill.solid(); r.fill.fore_color.rgb = RGBColor.from_string(fill)
    r.line.color.rgb = EDGE; r.line.width = Pt(0.75)
    r.shadow.inherit = False

def line_text(txt, left_px, top_px, size_px, name):
    """Place one text line so its ink starts at (left_px, top_px), as measured on the source image."""
    f = ImageFont.truetype(GILL, round(size_px), index=0)
    x0, y0, _, _ = f.getbbox(txt)
    tb = s.shapes.add_textbox(P(left_px - x0), P(top_px - y0), P(f.getlength(txt) + 8), P(size_px * 1.3))
    tb.name = name
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.TOP
    run = tf.paragraphs[0].add_run(); run.text = txt
    run.font.name = FONT; run.font.size = Pt(size_px * K); run.font.color.rgb = RGBColor(0, 0, 0)
    return tb

for txt, x, y in [("(4) arbitrary analysis", 418, 44),
                  ("(3) flexible playbooks + operators,", 229, 95), ("fixed DSL", 255, 118),
                  ("(2) flexible playbooks,", 146, 153), ("fixed operators + DSL", 172, 177),
                  ("(1) fixed playbooks", 72, 221)]:
    line_text(txt, x, y, LABEL_PX, "label " + txt)
line_text("more explainable", 70, 296, AXIS_PX, "axis left")
line_text("more expressive", 428, 296, AXIS_PX, "axis right")

# two-headed axis arrow
c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, P(55), P(276), P(595), P(276))
c.name = "axis arrow"
c.line.color.rgb = RGBColor(0, 0, 0); c.line.width = Pt(1)
ln = c.line._get_or_add_ln()
for tag in ("a:headEnd", "a:tailEnd"):
    e = etree.SubElement(ln, qn(tag)); e.set("type", "triangle"); e.set("w", "med"); e.set("len", "med")

prs.save("freedom-spectrum.pptx")
print("wrote freedom-spectrum.pptx")
